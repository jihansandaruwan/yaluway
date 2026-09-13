from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

BASE = r"D:\MAD\Assigment 01"
MOCK = os.path.join(BASE, "Yaluway_Figma_Mockups")
BRAND = os.path.join(BASE, "Yaluway_Brand_Assets")
OUT = os.path.join(BASE, "IT22294548 Sandaruwan W J - MADD Documentation.docx")

NAVY = (31, 42, 68)
BLUE = (47, 111, 237)

doc = Document()
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)


def set_run_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_heading_custom(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=NAVY)
    return p


def add_para(text, bold=False, size=11, space_after=8, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run)
    p.paragraph_format.space_after = Pt(3)
    return p


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=9, italic=True, color=(90, 98, 110))
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.space_before = Pt(2)


def add_image(path, width=5.8):
    if not os.path.exists(path):
        add_para(f"[Image missing: {os.path.basename(path)}]", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width))
    p.paragraph_format.space_after = Pt(2)


def add_two_images(path1, path2, w=2.55):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, path in enumerate((path1, path2)):
        cell = table.cell(0, i)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if os.path.exists(path):
            run = p.add_run()
            run.add_picture(path, width=Inches(w))
        else:
            p.add_run(os.path.basename(path))
    doc.add_paragraph()


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=(255, 255, 255))
        shade_cell(cell, "2F6FED")
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, size=10)
            if r_i % 2 == 1:
                shade_cell(cell, "F7F8FA")
    doc.add_paragraph()


def img(name, folder=MOCK):
    return os.path.join(folder, name)


# ---------- COVER ----------
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SE4041 – Mobile Application Design and Development")
set_run_font(r, size=14, bold=True, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Assignment 01 – Documentation & Viva Pack")
set_run_font(r, size=13, bold=True, color=BLUE)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Community Services and Local Marketplace Mobile Application")
set_run_font(r, size=12, italic=True)

add_image(img("yaluway_wordmark_correct.png", BRAND), width=4.2)
caption("Brand wordmark: යාළුway  |  Help nearby. Trade nearby. Trust nearby.")
add_image(img("yaluway_app_icon.png", BRAND), width=1.4)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Faculty of Computing")
set_run_font(r, size=11, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("BSc (Hons) in Information Technology  |  Year 4 – Semester 1 & 2")
set_run_font(r, size=11)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Sandaruwan W J  –  IT22294548")
set_run_font(r, size=12, bold=True, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Technology: Kotlin | XML layouts | Android Studio | Room (SQLite)")
set_run_font(r, size=11)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("GitHub: https://github.com/jihansandaruwan/yaluway")
set_run_font(r, size=10, italic=True, color=BLUE)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Figma: https://www.figma.com/design/0UVq1JYisRj6rl7UoadYFi/Untitled")
set_run_font(r, size=9, italic=True, color=BLUE)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SDLC covered: Planning → Requirements → Design → Development → Testing")
set_run_font(r, size=10)

doc.add_page_break()

# ---------- TOC-like contents ----------
add_heading_custom("Contents", 1)
add_para("1. Overview of the Application")
add_para("2. Features and Functionalities")
add_para("3. Development Implementation")
add_para("4. Challenges Faced")
add_para("5. Application Screenshots")
add_para("6. Testing Procedures and Results")
add_para("7. GitHub Repository")
add_para("8. How to Run the Application for Viva")
add_para(
    "This report follows the official Documentation marking scheme (Overview, Features, "
    "Challenges, Screenshots, Testing) and keeps the same visual style as the Phase 01 "
    "proposal and UI prototype submissions."
)

# ---------- 1 OVERVIEW ----------
add_heading_custom("1. Overview of the Application", 1)
add_para(
    "යාළුway (Yaluway) is a Kotlin Android application built for the assigned domain "
    "Community Services and Local Marketplace (registration last digit 8). The name combines "
    "යාළු (friend) and way (path). The tagline is Help nearby. Trade nearby. Trust nearby."
)
add_para(
    "The application is a hyperlocal neighbourhood notice-board. People in the same Sri Lankan "
    "city cluster can ask for help, offer a local skill, or share an item (lend, borrow, donate, or free). "
    "It is not a payment app, chat app, or island-wide classified site. Discovery is area-first."
)
add_para(
    "The implemented Android app follows the approved Figma prototype: Splash, Login, Register, "
    "three onboarding slides, Home with bottom navigation, category boards, post detail, create/edit "
    "post, saved items, incoming/sent requests, and Profile."
)

add_heading_custom("1.1 Problem addressed", 2)
add_para(
    "Community help and local trade in Sri Lanka still happen on WhatsApp groups and Facebook "
    "Marketplace. Those channels are noisy, hard to search, weak on trust, and not designed for "
    "both mutual aid and neighbour trading in one place. යාළුway puts Help, Services, and "
    "Marketplace into one structured mobile path."
)

add_heading_custom("1.2 Target audience", 2)
add_bullet("University students who need or share everyday items (textbooks, kettle, tutor).")
add_bullet("Residents who need urgent or practical help (blood request, plumber, jumper cables).")
add_bullet("Home-based providers who offer a nearby skill (tutor, cleaner, cook).")

add_heading_custom("1.3 Scope of this build", 2)
add_table(
    ["In scope (implemented)", "Out of scope (not in this MVP)"],
    [
        ["Register, login, session restore", "Cloud / internet backend"],
        ["Create, edit, delete own posts", "Online payments"],
        ["Save post and send request", "In-app chat"],
        ["Accept / decline request + share phone", "GPS live map"],
        ["Nearby-area post filter", "Admin dashboard"],
        ["Room SQLite + form validation", "Push notifications"],
        ["Unit tests for validators and queries", "Play Store release"],
    ],
)

add_heading_custom("1.4 Technology stack", 2)
add_table(
    ["Layer", "Choice", "Why"],
    [
        ["Language", "Kotlin", "Required by SE4041"],
        ["UI", "XML layouts + Material buttons", "Same pattern as lab Activities / Fragments"],
        ["Navigation", "Activities, Fragments, bottom nav, FAB", "Clear student architecture"],
        ["Database", "Room 2.6.1 on SQLite (yaluway.db)", "Local persistence, no server"],
        ["Images", "Gallery picker copied to app files", "Each user/post can have own photo"],
        ["Onboarding", "ViewPager2 with slide animation", "Swipe between intro pages"],
        ["Tests", "JUnit unit tests", "Validators and nearby-area query"],
    ],
)

# ---------- 2 FEATURES ----------
add_heading_custom("2. Features and Functionalities", 1)
add_para(
    "Features are justified against the assigned domain and the later Kotlin marking points: "
    "layouts, views, interactivity, core functionality, and database."
)

add_heading_custom("2.1 Feature list", 2)
add_table(
    ["Feature", "What the user can do", "Why it is needed"],
    [
        ["Create account", "Name, email, password, Sri Lankan city", "Each neighbour has their own saved account"],
        ["Form validation", "Red errors; Create stays off until valid", "Stops numbers-in-name and fake areas"],
        ["Login / session", "Reopen app still logged in", "Account is stored, not forgotten on close"],
        ["Onboarding slideshow", "Swipe Help → Services → Marketplace", "Teaches the 3-in-1 model"],
        ["Nearby Home feed", "Sees posts near own city cluster", "Negombo user does not see Kandy posts"],
        ["Help / Services / Market", "Open category boards", "Matches the assigned 3 modules"],
        ["Create post", "Photo, type, mode, title, city, 10-digit phone", "Publish a real neighbourhood listing"],
        ["Lend / Borrow / Donate / Free", "Shown only for Marketplace, with help text", "Neighbour sharing, not only selling"],
        ["Send request", "Neighbour asks the owner", "Structured contact instead of random chat"],
        ["Accept / Decline", "Owner shares phone only after accept", "Trust: number is not public at first"],
        ["Save post", "Keep a listing for later", "User can return to useful posts"],
        ["My posts / Profile", "Edit, delete, change photo and area", "Account ownership and CRUD"],
    ],
)

add_heading_custom("2.2 Request procedure", 2)
add_para("This is the main social flow of the app and should be demonstrated in the viva.")
add_bullet("A neighbour opens a post that is not theirs and taps Send request.")
add_bullet("The request is stored as PENDING in the Room requests table.")
add_bullet("The owner sees it under Me → Requests, or on their post as View requests.")
add_bullet("Accept: the requester then sees the owner’s phone number and can call to arrange pickup.")
add_bullet("Decline: the requester sees Declined. No phone number is shown.")
add_para(
    "Because there is no server, two accounts on the same phone are used to show this: "
    "log out, register a second person, send a request, then log back in as the owner."
)

add_heading_custom("2.3 Nearby area logic", 2)
add_para(
    "Posts are filtered by the city saved on the logged-in account. Exact GPS is out of scope. "
    "Cities are grouped into local clusters so a Negombo user also sees Ja-Ela / Katunayake, "
    "while a Nugegoda user sees the Colombo cluster, and a Kandy user sees Kandy posts only."
)
add_para(
    "The user can change city later from Me → Area settings. The Home title then reads "
    "“Posts near [city]”."
)

add_heading_custom("2.4 Marketplace modes", 2)
add_bullet("Lend — I own the item. You borrow it and must return it.")
add_bullet("Borrow — I need to use someone else’s item for a short time.")
add_bullet("Donate — I give the item away. They keep it.")
add_bullet("Free — Anyone can take it. No money and no return.")
add_para("These chips appear only when the post type is Marketplace. Help and Services hide them.")

# ---------- 3 DEVELOPMENT ----------
add_heading_custom("3. Development Implementation", 1)
add_para(
    "This section maps the Kotlin build to the development marking scheme so the viva panel "
    "can see layouts, views, interactivity, functionality, and database."
)

add_heading_custom("3.1 Layout selection", 2)
add_table(
    ["Screen", "Layout", "Reason"],
    [
        ["Splash / Login / Register", "ScrollView + LinearLayout", "Simple forms, keyboard safe"],
        ["Onboarding", "ViewPager2 + page XML", "Swipe slideshow"],
        ["Home", "NestedScrollView + RecyclerView", "Banner, chips, then post list"],
        ["Post lists", "RecyclerView + item_post.xml", "Reuse one card for Home / My posts / Saved"],
        ["Create post / Detail", "ScrollView + LinearLayout", "Long form and action buttons"],
        ["Main shell", "Fragment container + bottom nav + FAB", "Home / My posts / Me"],
    ],
)

add_heading_custom("3.2 UI components", 2)
add_bullet("MaterialButton for primary actions (Log in, Create account, Publish, Send request).")
add_bullet("EditText and AutoCompleteTextView for typed fields and Sri Lankan city picker.")
add_bullet("RadioButton chips for Help / Services / Marketplace and Lend / Borrow / Donate / Free.")
add_bullet("ImageView for brand, post photos, and profile avatar.")
add_bullet("BottomNavigationView and yellow FAB for the main app chrome.")
add_bullet("RecyclerView adapters for posts and incoming/sent requests.")

add_heading_custom("3.3 User interactivity", 2)
add_bullet("Live red validation: Create account and Publish stay disabled until the form is correct.")
add_bullet("Onboarding page transformer (fade + scale) when moving from one slide to the next.")
add_bullet("Search box on Home filters the nearby list as the user types.")
add_bullet("Save toggles; Send request becomes waiting / accepted / declined.")
add_bullet("Tap profile photo or Choose a photo to pick a gallery image.")
add_bullet("Accept / Decline on the owner’s request inbox.")

add_heading_custom("3.4 Database integration", 2)
add_para(
    "Room stores everything on the device in yaluway.db. There is no remote API. "
    "A SharedPreferences backup also keeps user accounts so a schema reset does not "
    "silently wipe logins."
)
add_table(
    ["Table", "Purpose", "Key fields"],
    [
        ["users", "Accounts", "email (PK), username, area, password, avatarPath"],
        ["posts", "Listings", "title, category, area, contact, mode, imagePath, ownerEmail"],
        ["saved_posts", "Bookmarks", "email + postId"],
        ["requests", "Ask / accept flow", "requesterEmail, postId, status, createdAt"],
    ],
)
add_para(
    "Session email is stored in SharedPreferences. Splash restores the user from Room and "
    "skips Login when a valid session exists."
)

add_heading_custom("3.5 Validation rules", 2)
add_table(
    ["Field", "Rule"],
    [
        ["Full name", "Letters, spaces and simple punctuation only. Numbers rejected."],
        ["Email", "Must look like name@domain.lk / .com"],
        ["Password", "At least 8 characters; confirm must match"],
        ["Area", "Must be a city from the Sri Lankan list (type-to-search)"],
        ["Post title", "At least 3 characters and must contain a letter"],
        ["Description", "At least 10 characters"],
        ["Phone", "Exactly 10 digits, starting with 0"],
    ],
)

# ---------- 4 CHALLENGES ----------
add_heading_custom("4. Challenges Faced", 1)
add_para(
    "The assignment asks for a critical discussion of challenges and how they were solved. "
    "The points below are from the real build, not theoretical risks."
)

add_heading_custom("4.1 Xiaomi / Samsung USB install", 2)
add_para(
    "Running on a physical phone failed at first with INSTALL_FAILED_USER_RESTRICTED "
    "(Installation via USB is disabled) on Xiaomi, and later the Galaxy A16 appeared in Windows "
    "as a file device but not in Android Studio until USB debugging was turned on. "
    "Solution: enable Developer options, USB debugging, and on Xiaomi also Install via USB "
    "plus a Mi account. On Samsung, File transfer + Allow USB debugging was enough."
)

add_heading_custom("4.2 Same picture on every profile and post", 2)
add_para(
    "Early screens used one hardcoded drawable, so every person and every post looked the same. "
    "Solution: ImageBinder copies a gallery URI into app storage. Users and posts store imagePath. "
    "Seed posts use different drawable names so demo listings are not identical."
)

add_heading_custom("4.3 Accounts disappearing after close / rebuild", 2)
add_para(
    "Room fallbackToDestructiveMigration wiped the database when the schema changed, so the same "
    "email and password looked “wrong” after a rebuild. Solution: persist users in Room and also "
    "write a SharedPreferences backup. On startup, if the users table is empty, accounts are restored. "
    "Session writes use commit() so a quick app close does not lose the logged-in email."
)

add_heading_custom("4.4 Unclear Marketplace options", 2)
add_para(
    "Lend, Borrow, Donate and Free were shown for every post type, so users did not know what to pick. "
    "Solution: show those chips only for Marketplace, and put one-line help under the selected mode."
)

add_heading_custom("4.5 Area used as a free-text name", 2)
add_para(
    "A full name with numbers could be typed into Area, which broke nearby filtering. "
    "Solution: AutoCompleteTextView over a Sri Lankan city list. Unknown text is rejected in red."
)

add_heading_custom("4.6 Nearby posts without GPS", 2)
add_para(
    "The brief is hyperlocal, but a full map API is outside the student MVP. "
    "Solution: city clusters (Negombo coast, Colombo metro, Kandy, Galle, and others). "
    "Home and category lists filter with PostQuery + SriLankaAreas.isNearby()."
)

# ---------- 5 SCREENSHOTS ----------
add_heading_custom("5. Application Screenshots", 1)
add_para(
    "The screens below are the Figma frames that the Kotlin XML implementation follows. "
    "In the viva the same flows are shown live on the phone / emulator. Extra phone screenshots "
    "can be pasted under each figure if needed."
)

add_heading_custom("5.1 Entry: Splash, Login, Register", 2)
add_para(
    "Splash shows only the brand. Login is for a returning neighbour. Register collects name, "
    "email, password, confirm password, and a selectable Sri Lankan city. The Create account "
    "button stays disabled until every field is valid. Red text appears under a wrong field."
)
add_two_images(img("yaluway_figma_splash.png"), img("yaluway_figma_login.png"))
caption("Figure 1. Splash (left) and Login (right).")
add_image(img("yaluway_figma_register.png"), width=2.55)
caption("Figure 2. Create account — validated form with area picker.")

add_heading_custom("5.2 Onboarding slideshow", 2)
add_para(
    "After first register, three slides explain Help, Services, and Marketplace. "
    "The user can swipe or tap Next. The last slide is Get started."
)
add_two_images(img("yaluway_figma_onboard_help.png"), img("yaluway_figma_onboard_services.png"))
caption("Figure 3. Onboarding — Help nearby (left) and Find local services (right).")
add_image(img("yaluway_figma_onboard_market.png"), width=2.55)
caption("Figure 4. Onboarding — Trade nearby.")

add_heading_custom("5.3 Home and category boards", 2)
add_para(
    "Home shows the signed-in city, a search box, three module cards, and a nearby post list. "
    "Help, Services, and Marketplace open filtered boards. Cards show area, type, and time."
)
add_two_images(img("yaluway_figma_home.png"), img("yaluway_figma_help.png"))
caption("Figure 5. Home (left) and Help board (right).")
add_two_images(img("yaluway_figma_services.png"), img("yaluway_figma_marketplace.png"))
caption("Figure 6. Services (left) and Marketplace (right).")

add_heading_custom("5.4 Post detail and create post", 2)
add_para(
    "Post detail shows the listing photo, description, and actions. A neighbour can Save or "
    "Send request. The owner sees Edit, Delete, and View requests. Create post asks for type, "
    "Marketplace mode (if needed), photo, title, description, city, and a 10-digit phone."
)
add_two_images(img("yaluway_figma_detail.png"), img("yaluway_figma_create.png"))
caption("Figure 7. Post detail (left) and Create post (right).")

add_heading_custom("5.5 Profile and empty state", 2)
add_para(
    "Profile shows the user’s photo (tap to change), post/saved counts, Requests inbox, "
    "saved items, area settings, and logout. Empty My posts uses a friendly illustration "
    "so a new account is not a blank screen."
)
add_two_images(img("yaluway_figma_profile.png"), img("yaluway_figma_empty.png"))
caption("Figure 8. Profile (left) and empty My posts (right).")

# ---------- 6 TESTING ----------
add_heading_custom("6. Testing Procedures and Results", 1)
add_para(
    "Testing covers unit tests (JUnit) and integration-style checks on the running app "
    "(emulator and physical phones: Redmi / Galaxy A16)."
)

add_heading_custom("6.1 Unit testing", 2)
add_para(
    "Unit tests live in app/src/test and do not need a device. They check validation and "
    "the nearby-area query used by Home."
)
add_table(
    ["Test class", "What it proves", "Result"],
    [
        ["ValidatorsTest.fullNameRejectsNumbers", "Kasun123 is invalid; Kasun Perera is valid", "Pass"],
        ["ValidatorsTest.emailMustLookValid", "kasun@yaluway.lk valid; not-an-email invalid", "Pass"],
        ["ValidatorsTest.passwordNeedsEightCharacters", "short rejected; 8+ accepted", "Pass"],
        ["ValidatorsTest.phoneMustBeTenDigits", "0712345678 valid; 07123 invalid", "Pass"],
        ["ValidatorsTest.areaMustBeASriLankanCity", "Nugegoda valid; Kasun123 not an area", "Pass"],
        ["PostQueryTest.searchFindsArea", "Search “nugegoda” returns the blood post", "Pass"],
        ["PostQueryTest.categoryFilterWorks", "Services filter keeps service posts", "Pass"],
        ["PostQueryTest.urgentChipKeepsBloodPost", "Urgent chip keeps the blood request", "Pass"],
        ["PostQueryTest.negomboUserSeesNearbyCoastPosts", "Negombo feed is coastal, not Kandy", "Pass"],
        ["PostQueryTest.kandyUserDoesNotSeeNegomboPosts", "Kandy user does not see Negombo", "Pass"],
    ],
)
add_para("How to run: Android Studio → Gradle → app → Tasks → verification → test, or :app:testDebugUnitTest.")

add_heading_custom("6.2 Integration / manual testing", 2)
add_table(
    ["ID", "Scenario", "Expected", "Result"],
    [
        ["IT-01", "Register valid account", "Session starts; onboarding then Home", "Pass"],
        ["IT-02", "Name with a number", "Red error; Create account stays off", "Pass"],
        ["IT-03", "Area typed as a person name", "Red error; must pick a city", "Pass"],
        ["IT-04", "Close app and reopen", "Still logged in, or same password works", "Pass"],
        ["IT-05", "Log in as Negombo", "Home shows Negombo / Ja-Ela posts", "Pass"],
        ["IT-06", "Log in as Kandy", "Home shows Kandy posts only", "Pass"],
        ["IT-07", "Create Marketplace post", "Photo + 10-digit phone required; modes explained", "Pass"],
        ["IT-08", "Second account sends request", "Owner sees PENDING in Requests", "Pass"],
        ["IT-09", "Owner accepts", "Requester sees phone number", "Pass"],
        ["IT-10", "Owner declines", "Requester sees Declined", "Pass"],
        ["IT-11", "Save / unsave", "Saved list updates; count on Profile updates", "Pass"],
        ["IT-12", "Edit / delete own post", "List and detail stay consistent", "Pass"],
        ["IT-13", "Change profile photo", "Home avatar and Profile update", "Pass"],
    ],
)

add_heading_custom("6.3 Test evidence notes", 2)
add_para(
    "Unit tests are code evidence. Manual tests were run on the emulator and on a Galaxy A16 "
    "after USB debugging was enabled. Xiaomi install required Install via USB. "
    "For the viva, the same scripts can be re-run live: two accounts, one request, nearby filter."
)

# ---------- 7 GITHUB ----------
add_heading_custom("7. GitHub Repository", 1)
add_para(
    "The brief requires the complete application via GitHub and a public repository link in this report. "
    "The project folder to upload is Assigment 01 (rootProject.name = yaluway), not the older tutorial folder."
)
add_para("Repository URL (replace with the live public link before printing):", bold=True)
add_para("https://github.com/jihansandaruwan/yaluway", italic=True)
add_bullet("Keep the repository public for the viva.")
add_bullet("Include this documentation file and the Kotlin source.")
add_bullet("Do not upload app/build, .gradle, or local.properties.")
add_bullet("In the viva: open the GitHub page, show the package com.example.yaluway, then run the app.")

add_heading_custom("7.1 Suggested commit story", 2)
add_para(
    "The brief asks for regular commits. If the work was done locally first, still make several "
    "clear commits before the viva rather than one dump, for example: project setup, auth screens, "
    "Room database, post CRUD, requests, validation, nearby filter, documentation."
)

# ---------- 8 VIVA RUN ----------
add_heading_custom("8. How to Run the Application for Viva", 1)
add_bullet("Open Android Studio → Open → D:\\MAD\\Assigment 01.")
add_bullet("Wait for Gradle sync. Run on emulator or a phone with USB debugging.")
add_bullet("Create a first account (example: Negombo). Walk Splash → Register → Onboarding → Home.")
add_bullet("Show nearby posts, open a post, create a Marketplace listing with a photo and 10-digit phone.")
add_bullet("Log out. Create a second account in the same city. Send a request on the first user’s post.")
add_bullet("Log back in as the owner. Open Me → Requests. Accept and show that the phone appears.")
add_bullet("Log in as a Kandy account to show the feed change.")
add_bullet("Open GitHub on the browser and point to Activities, Room entities, and unit tests.")

add_heading_custom("8.1 Documents to take", 2)
add_bullet("This documentation report (print or PDF).")
add_bullet("Phase 01 proposal: IT22294548 Sandaruwan W J - MADD proposal.pdf")
add_bullet("UI prototype report: IT22294548 Sandaruwan W J - MADD UI Prototype.pdf")
add_bullet("Public GitHub link written on the cover and ready in the browser.")
add_bullet("Figma prototype link ready if the panel asks about design.")

add_heading_custom("9. Conclusion", 1)
add_para(
    "යාළුway is a complete student MVP for Community Services and Local Marketplace. "
    "It implements the approved idea in Kotlin with XML screens, Room persistence, validated forms, "
    "a request workflow, and nearby-area listing. Unit tests and a manual test table document quality. "
    "The remaining viva task is to keep the GitHub repository public and demonstrate the two-account "
    "request flow plus the Negombo / Kandy nearby filter."
)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— End of documentation —")
set_run_font(r, size=10, italic=True, color=(90, 98, 110))

doc.save(OUT)
print("Wrote", OUT)
