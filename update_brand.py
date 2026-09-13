from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

out = r"D:\MAD\Assigment 01\LankaWay_Phase01_Proposal.docx"
doc = Document(out)


def set_run_font(run, size=11, bold=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold


for p in doc.paragraphs:
    if "Help nearby. Trade nearby. Trust nearby." in p.text and "LankaWay aims" in p.text:
        for r in p.runs:
            r.text = ""
        if p.runs:
            p.runs[0].text = (
                'Brand meaning: "Lanka" represents Sri Lanka / local community, and "Way" (maga / path) '
                "means the path to get help, offer services, and trade nearby. "
                'Tagline: "Help nearby. Trade nearby. Trust nearby." '
                "LankaWay aims to become the go-to mobile hub for neighborhood support and "
                "small-scale local exchange, combining mutual aid with a structured local marketplace."
            )
            set_run_font(p.runs[0])
        break

for p in doc.paragraphs:
    if "A Hyperlocal Community Services" in p.text:
        for r in p.runs:
            r.text = ""
        if p.runs:
            p.runs[0].text = (
                "A Hyperlocal Community Services &\n"
                "Local Marketplace Mobile Application\n"
                "(Sinhala-English inspired brand name)"
            )
            set_run_font(p.runs[0], size=12)
        break

doc.save(out)
print("Updated:", out)
