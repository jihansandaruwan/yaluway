from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
import os

src = r"D:\MAD\Assigment 01\LankaWay_Phase01_Proposal.docx"
if not os.path.exists(src):
    src = r"D:\MAD\Assigment 01\NeighbrLink_Phase01_Proposal.docx"

out = r"D:\MAD\Assigment 01\Yaluway_Phase01_Proposal.docx"
doc = Document(src)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


replacements = 0
for p in doc.paragraphs:
    for run in p.runs:
        text = run.text
        if not text:
            continue
        new = text.replace("LankaWay", "Yaluway").replace("NeighbrLink", "Yaluway")
        if new != text:
            run.text = new
            replacements += 1

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    text = run.text
                    if not text:
                        continue
                    new = text.replace("LankaWay", "Yaluway").replace(
                        "NeighbrLink", "Yaluway"
                    )
                    if new != text:
                        run.text = new
                        replacements += 1

# Refresh brand meaning / vision paragraph
for p in doc.paragraphs:
    if "Brand meaning:" in p.text or (
        "Help nearby. Trade nearby. Trust nearby." in p.text and "Yaluway aims" in p.text
    ):
        for r in p.runs:
            r.text = ""
        if p.runs:
            p.runs[0].text = (
                'Brand meaning: "Yalu" (from Sinhala yaluwa / friend) + "Way" (path). '
                "Together, Yaluway means a friendly local path to get help, offer services, "
                "and trade with people nearby. "
                'Tagline: "Help nearby. Trade nearby. Trust nearby." '
                "Yaluway aims to become the go-to mobile hub for neighborhood support and "
                "small-scale local exchange, combining mutual aid with a structured local marketplace."
            )
            set_run_font(p.runs[0])
        break

# Cover subtitle note
for p in doc.paragraphs:
    if "A Hyperlocal Community Services" in p.text:
        for r in p.runs:
            r.text = ""
        if p.runs:
            p.runs[0].text = (
                "A Hyperlocal Community Services &\n"
                "Local Marketplace Mobile Application\n"
                '(Sinhala-English brand: Yalu = friend + Way = path)'
            )
            set_run_font(p.runs[0], size=12)
        break

# Ensure cover title run is styled if it is just "Yaluway"
for p in doc.paragraphs:
    if p.text.strip() == "Yaluway":
        for r in p.runs:
            set_run_font(r, size=28, bold=True, color=(20, 90, 70))

# Appendix selection sentence if it still says LankaWay was selected
for p in doc.paragraphs:
    if "was selected because" in p.text:
        for r in p.runs:
            r.text = ""
        if p.runs:
            p.runs[0].text = (
                "Yaluway (same concept as the original NeighbrLink idea) was selected "
                "because it best balances innovation, market gap, topic coverage, and MVP feasibility."
            )
            set_run_font(p.runs[0])
        break

doc.save(out)
print("Source:", src)
print("Saved:", out)
print("Run replacements:", replacements)

# Verify
d2 = Document(out)
has_old = False
yalu_count = 0
for p in d2.paragraphs:
    if "LankaWay" in p.text or "NeighbrLink" in p.text:
        # NeighbrLink may remain once in appendix on purpose
        if "NeighbrLink" in p.text and "same concept" in p.text:
            continue
        has_old = True
        print("OLD LEFT:", p.text[:120])
    if "Yaluway" in p.text:
        yalu_count += 1
print("Yaluway paragraphs:", yalu_count)
print("Unexpected old names left:", has_old)
