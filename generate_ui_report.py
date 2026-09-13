from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
import os

BASE = r"D:\MAD\Assigment 01"
MOCK = os.path.join(BASE, "Yaluway_Figma_Mockups")
BRAND = os.path.join(BASE, "Yaluway_Brand_Assets")
OUT = os.path.join(BASE, "IT22294548 Sandaruwan W J - MADD UI Prototype.docx")

NAVY = (31, 42, 68)
BLUE = (47, 111, 237)
YELLOW = (245, 197, 24)

doc = Document()
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)


def set_run_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D0D5DD")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_heading_custom(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=NAVY)
    return p


def add_para(text, bold=False, size=11, space_after=8, italic=False, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if align:
        p.alignment = align
    return p


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run)
    p.paragraph_format.space_after = Pt(3)
    return p


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=9, italic=True, color=(90, 98, 110))
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.space_before = Pt(2)
    return p


def add_image(path, width=5.8):
    if not os.path.exists(path):
        add_para(f"[Image missing: {os.path.basename(path)}]", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width))
    p.paragraph_format.space_after = Pt(2)


def add_two_images(path1, path2, w=2.55):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, path in enumerate((path1, path2)):
        cell = table.cell(0, i)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if os.path.exists(path):
            run = p.add_run()
            run.add_picture(path, width=Inches(w))
        else:
            p.add_run(os.path.basename(path))
    doc.add_paragraph()


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=(255, 255, 255))
        shade_cell(cell, "2F6FED")
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, size=10)
            if r_i % 2 == 1:
                shade_cell(cell, "F7F8FA")
    doc.add_paragraph()


def img(name, folder=MOCK):
    return os.path.join(folder, name)


# ---------- COVER ----------
for _ in range(1):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SE4041 – Mobile Application Design and Development")
set_run_font(r, size=14, bold=True, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Assignment 01 – UI Prototype Submission")
set_run_font(r, size=13, bold=True, color=BLUE)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Community Services and Local Marketplace Mobile Application")
set_run_font(r, size=12, italic=True)

add_image(img("yaluway_wordmark_correct.png", BRAND), width=4.2)
caption("Brand wordmark: යාළුway  |  Help nearby. Trade nearby. Trust nearby.")

add_image(img("yaluway_app_icon.png", BRAND), width=1.5)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Faculty of Computing")
set_run_font(r, size=11, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("BSc (Hons) in Information Technology  |  Year 4 – Semester 1 & 2")
set_run_font(r, size=11)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Sandaruwan W J  –  IT22294548")
set_run_font(r, size=12, bold=True, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Figma prototype: https://www.figma.com/design/0UVq1JYisRj6rl7UoadYFi/Untitled")
set_run_font(r, size=9, italic=True, color=BLUE)

doc.add_page_break()

# ---------- 1. INTRODUCTION ----------
add_heading_custom("1. Introduction", 1)
add_para(
    "This document submits the Figma UI prototype for යාළුway (Yaluway), designed under the assigned "
    "domain Community Services and Local Marketplace. The prototype covers ideation, the 60-30-10 colour "
    "rule, main interface design, and aesthetics and usability, matching the SE4041 marking scheme."
)
add_para(
    "යාළුway is a hyperlocal neighbourhood notice-board app. People nearby can post three kinds of things "
    "in one place: community Help, local Services, and Marketplace goods (sell, donate, or borrow). "
    "The UI is designed as one listing app with three post types, not three separate products."
)

add_heading_custom("1.1 Purpose of this submission", 2)
add_bullet("Present the complete Figma prototype screens with labelled explanations.")
add_bullet("Show how ideation is translated into interface structure and visual identity.")
add_bullet("Demonstrate consistent 60-30-10 colour application across all screens.")
add_bullet("Explain usability choices that support the later Kotlin / Room implementation.")

add_heading_custom("1.2 Prototype overview", 2)
add_para(
    "The prototype contains 15 mobile frames (iPhone 14 / 390 × 844) plus a Design system page. "
    "Screens are linked in Figma Prototype so the flow can be clicked from Splash to Profile."
)
add_image(img("yaluway_figma_overview.png"), width=6.4)
caption("Figure 1. Complete Figma UI prototype (all screens on one canvas).")

# ---------- 2. IDEATION ----------
add_heading_custom("2. Ideation", 1)
add_para(
    "Ideation marks assess whether the idea is innovative, fits the assigned domain, and clearly "
    "addresses user needs. The UI is the visual proof of that idea."
)

add_heading_custom("2.1 Problem the interface solves", 2)
add_para(
    "In Sri Lanka, neighbourhood help and local trade still happen on WhatsApp groups and Facebook "
    "Marketplace. Those channels are noisy, hard to search, weak on trust, and not designed for both "
    "community service and local commerce together. යාළුway’s interface puts Help, Services, and "
    "Marketplace in one structured mobile path."
)

add_heading_custom("2.2 Brand meaning used in the UI", 2)
add_para(
    "Yalu (යාළු / friend) + Way (path) means a friendly local path to get help, offer services, "
    "and trade nearby. Tagline: Help nearby. Trade nearby. Trust nearby."
)
add_bullet("Yellow in the wordmark stands for friendship and warmth (යාළු).")
add_bullet("Blue in the wordmark stands for trust and direction (way).")
add_bullet("Handshake / pin icon represents two neighbours meeting on a local path.")

add_heading_custom("2.3 How ideation appears in the screens", 2)
add_table(
    ["UI decision", "Why it supports the idea"],
    [
        ["Home chips: Help | Services | Marketplace", "Shows the 3-in-1 model immediately."],
        ["Onboarding in three slides", "Teaches the three modules before first use."],
        ["Borrow / Donate / Sell chips", "Goes beyond classified ads; supports neighbour culture."],
        ["Area label (Nugegoda / Hostel A)", "Hyperlocal, not island-wide anonymous listings."],
        ["Trust score on Profile", "Neighbour trust is a core differentiator."],
        ["No payment or chat screens", "MVP stays academic and buildable in Kotlin + Room."],
    ],
)

add_heading_custom("2.4 Target users reflected in content", 2)
add_para(
    "Screen copy uses Sri Lankan neighbourhood examples so the panel can see the audience: "
    "students (textbooks, kettle, maths tutor), residents (plumber, blood request), and home "
    "providers (baker, cleaner). Persona Nimali Perera appears on Profile and post-detail."
)

# ---------- 3. 60-30-10 ----------
add_heading_custom("3. Application of the 60-30-10 Colour Rule", 1)
add_para(
    "The 60-30-10 rule controls how much of the screen each colour occupies. It is not three "
    "colours mixed equally. 60% is the dominant canvas, 30% is the primary brand/action colour, "
    "and 10% is a small accent used for attention."
)

add_heading_custom("3.1 Colour palette", 2)
add_image(img("yaluway_figma_ui_kit.png"), width=6.2)
caption("Figure 2. Design system / 60-30-10 palette used in the prototype.")

add_table(
    ["Share", "Role", "Hex", "Where it is used"],
    [
        ["60%", "Background", "#F7F8FA", "Screen canvas on every frame"],
        ["30%", "Primary (Way blue)", "#2F6FED", "Buttons, selected chips, links, “way”"],
        ["10%", "Accent (Yalu yellow)", "#F5C518", "FAB +, Borrow badge, යාළු wordmark"],
        ["Support", "Text navy", "#1F2A44", "Headings and body text"],
        ["Support", "Card white", "#FFFFFF", "Cards, input fields"],
        ["Support", "Success teal", "#2BBBAD", "Success / available states"],
    ],
)

add_heading_custom("3.2 How the rule is applied on a typical screen", 2)
add_para("Example — Home:")
add_bullet("60%: large off-white background around header, chips, and cards.")
add_bullet("30%: selected chip, primary actions, and blue brand elements.")
add_bullet("10%: yellow + FAB and yellow BORROW badge. Small area, high noticeability.")
add_para(
    "The same split is repeated on Login (blue Log in button, yellow Create account link), "
    "Marketplace (blue title/actions, yellow selected Borrow chip), and post-detail (blue Send request, "
    "yellow BORROW · FREE pill). Yellow is never used as a second background. If yellow covered half "
    "the screen, the 10% rule would break."
)

add_heading_custom("3.3 Consistency check", 2)
add_para(
    "Colour styles named Yaluway/Background, Yaluway/Primary Blue, Yaluway/Accent Yellow, "
    "Yaluway/Text Navy, Yaluway/Success Teal, and Yaluway/Cards were created in Figma so every "
    "screen uses the same tokens."
)

# ---------- 4. IA / FLOW ----------
add_heading_custom("4. Information Architecture and User Flow", 1)
add_para(
    "Navigation is kept simple so the same pattern can be coded later with Activities / Fragments "
    "and a bottom bar."
)

add_table(
    ["Stage", "Screens", "User goal"],
    [
        ["Entry", "Splash → Login / Register", "Recognise brand and sign in"],
        ["Learn", "Onboarding Help → Services → Marketplace", "Understand the 3 modules"],
        ["Discover", "Home, Help board, Services, Marketplace", "Browse nearby posts"],
        ["Act", "Post detail, Create post", "Request help or publish a listing"],
        ["Manage", "My posts, Empty state, Profile", "See own activity and account"],
    ],
)

add_para("Prototype click path used for demonstration:", bold=True)
add_para(
    "Splash → Login → Home. New user: Login Create account → Register → Onboarding Help → "
    "Onboarding Services → Onboarding Marketplace → Home. From Home: chips open Help / Services / "
    "Marketplace; a card opens Post detail; yellow + opens Create post; bottom nav opens My posts "
    "or Profile."
)

# ---------- 5. MAIN INTERFACES ----------
add_heading_custom("5. Main Interface Design", 1)
add_para(
    "This section presents each screen with a screenshot and a short explanation of purpose, "
    "layout, and usability. Screens are grouped by user journey."
)

# 5.1 splash login register
add_heading_custom("5.1 Splash, Login and Register", 2)
add_para(
    "Splash introduces the brand without clutter. Login is the launcher for returning users. "
    "Register collects only what the MVP needs: name, email, password, confirm password, and area. "
    "Area (example: Nugegoda) is required because discovery is neighbourhood-first."
)
add_two_images(img("yaluway_figma_splash.png"), img("yaluway_figma_login.png"))
caption("Figure 3. Splash (left) and Login (right).")
add_image(img("yaluway_figma_register.png"), width=2.55)
caption("Figure 4. Register — Create account.")
add_bullet("Primary action is always a full-width blue button (30% colour).")
add_bullet("Create account on Login is yellow so it is visible without competing with Log in.")
add_bullet("No bottom navigation on these screens — the user is not inside the app yet.")

# 5.2 onboarding
add_heading_custom("5.2 Onboarding (Help, Services, Marketplace)", 2)
add_para(
    "Three slides teach the 3-in-1 model before Home. Pagination dots show progress. "
    "Get started on the last slide enters the app. Illustrations use the same yellow-blue language "
    "as the logo so brand recognition starts before login content appears."
)
add_two_images(img("yaluway_figma_onboard_help.png"), img("yaluway_figma_onboard_services.png"))
caption("Figure 5. Onboarding 1 — Help nearby (left). Onboarding 2 — Find local services (right).")
add_image(img("yaluway_figma_onboard_market.png"), width=2.55)
caption("Figure 6. Onboarding 3 — Trade nearby, with Get started.")

# 5.3 home
add_heading_custom("5.3 Home — main dashboard", 2)
add_para(
    "Home is the main interface. It answers: where am I, what can I do, and what is nearby. "
    "Location (Nugegoda), search, three module chips, a mixed recent-posts feed, yellow FAB, "
    "and bottom navigation (Home | My posts | Me) are all on one screen."
)
add_image(img("yaluway_figma_home.png"), width=2.7)
caption("Figure 7. Home feed with Help / Services / Marketplace chips and yellow FAB.")
add_bullet("Chips filter the same list pattern — not three different apps.")
add_bullet("Card badges (HELP / SERVICE / BORROW) make post type scannable.")
add_bullet("FAB is yellow (10%) so “create” is always findable without a fifth nav item.")

# 5.4 help
add_heading_custom("5.4 Help board", 2)
add_para(
    "Help board is for urgent neighbourhood needs: blood requests, groceries for elders, "
    "volunteer tasks. A banner illustration sets the community-care tone. Filter chips "
    "(All, Urgent, Volunteer, Blood) reduce noise compared with a WhatsApp group."
)
add_image(img("yaluway_figma_help.png"), width=2.7)
caption("Figure 8. Community Help board.")

# 5.5 services
add_heading_custom("5.5 Services directory", 2)
add_para(
    "Services lists local skills: tutors, plumbers, cleaners, home cooks. Category cards "
    "with simple icons are faster than reading long classified ads. This screen covers the "
    "“community services” half of Topic 8 together with Help board."
)
add_image(img("yaluway_figma_services.png"), width=2.7)
caption("Figure 9. Local services directory.")

# 5.6 marketplace
add_heading_custom("5.6 Marketplace", 2)
add_para(
    "Marketplace uses a two-column grid because users compare goods visually. Sell / Donate / "
    "Borrow chips are the key differentiator from ikman-style ads. Borrow selected in yellow "
    "shows the 10% accent used for a meaningful mode, not decoration."
)
add_image(img("yaluway_figma_marketplace.png"), width=2.7)
caption("Figure 10. Marketplace grid with Sell / Donate / Borrow.")

# 5.7 detail create
add_heading_custom("5.7 Post detail and Create post", 2)
add_para(
    "Post detail shows one listing in full: hero image, type badge, title, area, poster, "
    "status, description, rating, Send request, and Save. Create post reuses one form for all "
    "three modules. Type (Help / Services / Marketplace) and Mode (Sell / Donate / Borrow) "
    "are chips, so the Kotlin app can store one Post table with type fields."
)
add_two_images(img("yaluway_figma_detail.png"), img("yaluway_figma_create.png"))
caption("Figure 11. Post detail (left) and Create post (right).")

# 5.8 empty my posts profile
add_heading_custom("5.8 Empty state, My posts and Profile", 2)
add_para(
    "Empty state is shown when My posts has no data — important for first-time users and for "
    "usability marks. The filled My posts list shows the user’s own listings. Profile holds "
    "identity, area, trust score, stats (posts / saved / rating), and account actions including Logout."
)
add_two_images(img("yaluway_figma_empty.png"), img("yaluway_figma_profile.png"))
caption("Figure 12. Empty state (left) and Profile (right).")

# ---------- 6. AESTHETICS ----------
add_heading_custom("6. Design Aesthetics and Usability", 1)

add_heading_custom("6.1 Visual aesthetics", 2)
add_bullet("Flat Material-like UI: no heavy shadows, no gradients, no extra purple.")
add_bullet("Consistent 12 px rounded cards and 48 px primary buttons for a professional look.")
add_bullet("Illustrations share one style (blue + yellow neighbourhood scenes) so screens feel like one product.")
add_bullet("Typography hierarchy: large navy titles, smaller body, caption metadata in muted grey.")
add_bullet("Plenty of whitespace (60% background) so the UI does not feel crowded like social feeds.")

add_heading_custom("6.2 Usability decisions", 2)
add_table(
    ["Principle", "How the prototype applies it"],
    [
        ["Clarity", "One primary button per screen; labels are verbs (Log in, Publish, Send request)."],
        ["Consistency", "Same header, FAB, and bottom nav pattern on all in-app screens."],
        ["Feedback", "Selected chips change fill colour; empty state explains what to do next."],
        ["Error prevention", "Register has Confirm password; Create post uses chips instead of free typing for type."],
        ["Learnability", "Onboarding + Home chips teach the three modules in under a minute."],
        ["Accessibility", "Navy text on off-white; large tap targets; high contrast blue buttons."],
        ["Efficiency", "FAB is always available to create a post without opening a menu."],
        ["Local context", "Area fields (Nugegoda, Hostel A) keep tasks neighbourhood-sized."],
    ],
)

add_heading_custom("6.3 What was deliberately left out", 2)
add_para(
    "Chat, payments, live maps, and AI matching are out of scope for the academic MVP. "
    "Hiding them from the prototype keeps the interface honest and matches the Kotlin + Room "
    "build plan. Users arrange pickup using profile contact outside the app."
)

# ---------- 7. COMPONENTS ----------
add_heading_custom("7. Reusable Components (Design System)", 1)
add_para(
    "A Design system page in Figma holds colour styles and components so screens stay consistent "
    "and later Android XML / Compose can map 1:1."
)
add_table(
    ["Component", "Variants / use"],
    [
        ["Primary button", "Filled #2F6FED, white text — Log in, Publish, Next"],
        ["Outline button", "Blue border — Save, secondary actions"],
        ["Input field", "Rounded, labelled — email, password, title, area"],
        ["Chip", "Unselected outline / selected fill — filters and modes"],
        ["Post card", "White card, badge, title, metadata"],
        ["Bottom navigation", "Home, My posts, Me with selected state"],
        ["FAB", "Yellow circle + — Create post"],
    ],
)

# ---------- 8. PROTOTYPE ----------
add_heading_custom("8. Interactive Prototype", 1)
add_para(
    "Figma Prototype connections let the panel click through the app. Starting point is Splash."
)
add_table(
    ["From", "Control", "Goes to"],
    [
        ["Splash", "Tap screen", "Login"],
        ["Login", "Log in", "Home"],
        ["Login", "Create account", "Register"],
        ["Register", "Create account", "Onboarding Help"],
        ["Onboarding Help", "Next", "Onboarding Services"],
        ["Onboarding Services", "Next", "Onboarding Marketplace"],
        ["Onboarding Marketplace", "Get started", "Home"],
        ["Home", "Help / Services / Marketplace chips", "Matching board"],
        ["Home", "Post card", "Post detail"],
        ["Any in-app screen", "Yellow + FAB", "Create post"],
        ["Create post", "Publish post", "Home"],
        ["Empty state", "Create first post", "Create post"],
        ["Bottom nav", "My posts / Me", "My posts or Profile"],
        ["Post detail", "Back", "Home"],
    ],
)
add_para(
    "Play the prototype in Figma: open the file → Prototype tab → Play (present). "
    "Link: https://www.figma.com/design/0UVq1JYisRj6rl7UoadYFi/Untitled"
)

# ---------- 9. MAPPING TO MARKS ----------
add_heading_custom("9. Mapping to the Marking Scheme", 1)
add_table(
    ["Criterion", "Max", "Evidence in this prototype"],
    [
        ["Ideation", "3", "3-in-1 hyperlocal Help + Services + Marketplace; brand යාළුway; Sri Lankan neighbourhood scenarios."],
        ["60-30-10 colour rule", "2", "Off-white 60 / blue 30 / yellow 10 applied on every screen; documented in Section 3."],
        ["Main interface design", "3", "15 frames covering auth, onboarding, home, 3 boards, detail, create, empty, my posts, profile."],
        ["Aesthetics and usability", "2", "Consistent components, clear hierarchy, chips, FAB, empty state, contrast, no clutter."],
    ],
)

# ---------- 10. CONCLUSION ----------
add_heading_custom("10. Conclusion", 1)
add_para(
    "The යාළුway Figma prototype is a complete, clickable UI for Topic 8. It turns a fragmented "
    "social-media problem into one neighbourhood notice board. Colour follows 60-30-10, screens "
    "share one component system, and the flow is ready to implement in Kotlin with Room. "
    "The next SDLC stage is development: the same screens become XML/Compose layouts, with "
    "posts stored locally and regular GitHub commits."
)

add_heading_custom("Appendix A — Screen inventory", 1)
add_table(
    ["#", "Frame name", "Purpose"],
    [
        ["1", "Splash", "Brand entry"],
        ["2", "Login", "Returning user authentication"],
        ["3", "Register", "New account + area"],
        ["4", "onboarding-help", "Explain Help module"],
        ["5", "onboarding-services", "Explain Services module"],
        ["6", "onboarding-marketplace", "Explain Marketplace module"],
        ["7", "Home", "Dashboard / mixed feed"],
        ["8", "help-board", "Community requests"],
        ["9", "Services", "Local service categories"],
        ["10", "Marketplace", "Goods grid Sell/Donate/Borrow"],
        ["11", "post-detail", "Single listing"],
        ["12", "create-post", "Publish a listing"],
        ["13", "empty-state", "No posts yet"],
        ["14", "my-posts", "User’s own listings"],
        ["15", "Profile", "Account, trust, settings"],
    ],
)

add_heading_custom("Appendix B — Figma file", 1)
add_para("File name: Untitled (recommend renaming to Yaluway UI Prototype before submit).")
add_para("URL: https://www.figma.com/design/0UVq1JYisRj6rl7UoadYFi/Untitled")
add_para(
    "Before uploading to Course Web: set the Figma file to “Anyone with the link can view”, "
    "paste the link in this report, export this document as PDF if the LMS requires PDF, "
    "and include your name and registration number on the cover (already filled)."
)

doc.save(OUT)
print("Saved:", OUT)
print("Size:", os.path.getsize(OUT))
