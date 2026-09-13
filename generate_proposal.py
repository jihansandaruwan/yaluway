from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_custom(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(
            run,
            size=16 if level == 1 else 13 if level == 2 else 12,
            bold=True,
        )
    return p


def add_para(text, bold=False, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullet(text):
    p = doc.add_paragraph(text, style="List Bullet")
    if p.runs:
        set_run_font(p.runs[0])
    return p


for _ in range(2):
    doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("SE4041 – Mobile Application Design and Development")
set_run_font(r, size=14, bold=True)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Assignment 01 – Phase 01 Proposal")
set_run_font(r, size=13, bold=True)

app = doc.add_paragraph()
app.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = app.add_run("\nNeighbrLink")
set_run_font(r, size=28, bold=True, color=(20, 90, 70))

tag = doc.add_paragraph()
tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tag.add_run(
    "A Hyperlocal Community Services &\nLocal Marketplace Mobile Application"
)
set_run_font(r, size=12)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run(
    "\nAssigned Topic (Last Digit 8):\n"
    "Community Services and Local Marketplace Mobile Application"
)
set_run_font(r, size=11)

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = info.add_run(
    "\n\nFaculty of Computing\n"
    "BSc (Hons) in Information Technology\n"
    "Year 4 – Semester 1 & 2\n\n"
    "Technology Stack: Kotlin | Android Studio | Local Database\n"
    "SDLC Focus: Planning → Requirements → Design → Development → Testing"
)
set_run_font(r, size=11)

# Student details placeholder
details = doc.add_paragraph()
details.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = details.add_run(
    "\n\nStudent Name: ______________________________\n"
    "Registration Number: ________________________\n"
    "Date: _____________________________________"
)
set_run_font(r, size=11)

doc.add_page_break()

add_heading_custom("1. Introduction & Proposed Idea", 1)
add_para(
    "This proposal presents NeighbrLink, a Kotlin-based Android mobile application "
    "designed under the assigned domain “Community Services and Local Marketplace.” "
    "The application connects people within a defined local area (neighborhood, "
    "apartment complex, university zone, or city ward) so they can request and offer "
    "community help, discover trusted local service providers, and buy, sell, or borrow "
    "goods safely."
)

add_heading_custom("1.1 Why This Idea?", 2)
add_para(
    "In Sri Lanka and many similar markets, community help and local trade still rely "
    "heavily on WhatsApp groups, Facebook Marketplace, and informal word-of-mouth. "
    "These channels are noisy, hard to search, weak on trust, and not designed for both "
    "community service coordination and local commerce in one place. NeighbrLink fills "
    "that gap with a purpose-built hyperlocal platform focused on trust, proximity, and "
    "practical everyday usefulness."
)

add_heading_custom("1.2 Application Vision", 2)
add_para(
    "“Help nearby. Trade nearby. Trust nearby.” NeighbrLink aims to become the go-to "
    "mobile hub for neighborhood support and small-scale local exchange, combining "
    "mutual aid with a structured local marketplace."
)

add_heading_custom("2. Problem Statement", 1)
add_para(
    "Current solutions create several practical problems for residents, students, and "
    "informal service workers:"
)
add_bullet(
    "Community needs (blood requests, volunteer help, lost items, elder support, tool "
    "borrowing) are scattered across chat groups and social media posts that expire quickly."
)
add_bullet(
    "Local buying/selling is dominated by general classified platforms that are city-wide, "
    "low-trust, and not community-oriented."
)
add_bullet(
    "Skilled informal workers (tutors, electricians, cleaners, home cooks) struggle to get "
    "visible, verified local demand without paying for ads or relying on random referrals."
)
add_bullet(
    "Users have no single place to request help, offer services, and trade goods within a "
    "trusted radius."
)
add_bullet(
    "Safety and credibility are weak: limited verification, weak ratings for community "
    "help, and poor category structure."
)
add_para(
    "NeighbrLink addresses these issues by providing one mobile application for community "
    "service posts, service-provider discovery, and a hyperlocal goods marketplace with "
    "profiles, categories, ratings, and local filters."
)

add_heading_custom("3. Market Gap Analysis", 1)
add_para(
    "Existing platforms partially cover this space, but none fully serve the hyperlocal "
    "community + marketplace combination for the Sri Lankan / South Asian neighborhood context:"
)
add_bullet(
    "Facebook Marketplace / Groups: High usage, but cluttered, unsafe messaging, weak "
    "service booking, and poor discovery for urgent community needs."
)
add_bullet(
    "ikman.lk / classified sites: Strong for ads, weak for community mutual aid, volunteer "
    "matching, and neighborhood-level trust."
)
add_bullet(
    "TaskRabbit / Airtasker-style apps: Useful abroad for task hiring, but limited local "
    "presence, high formality, and not focused on community volunteering or goods exchange."
)
add_bullet(
    "Donation / volunteer apps: Often single-purpose (blood, charity) and do not support "
    "local trade or micro-services."
)
add_para(
    "Market Gap: There is no widely adopted mobile application that unifies (1) community "
    "help requests, (2) verified local micro-services, and (3) neighborhood buy–sell–borrow "
    "trading with trust signals and location relevance. NeighbrLink targets exactly this "
    "underserved intersection."
)

add_heading_custom("3.1 Opportunity Justification", 2)
add_bullet("Rising smartphone penetration and preference for mobile-first local discovery.")
add_bullet(
    "Strong cultural habit of neighborhood mutual help that is currently digitally unstructured."
)
add_bullet("Growth of home-based and informal services that need low-friction visibility.")
add_bullet(
    "Student and apartment communities need safer, localized exchange than open public marketplaces."
)

add_heading_custom("4. Target Audience", 1)
add_heading_custom("4.1 Primary Users", 2)
add_bullet("Residents of neighborhoods, housing schemes, and apartment complexes (age 18–45).")
add_bullet(
    "University students living in hostels / boarding places who buy, sell, borrow, and need local help."
)
add_bullet(
    "Informal and semi-formal local service providers (tutors, repair persons, cleaners, bakers, caregivers)."
)

add_heading_custom("4.2 Secondary Users", 2)
add_bullet("Community admins / apartment managers who moderate local zones.")
add_bullet("Volunteer groups and donation organizers coordinating short local campaigns.")
add_bullet("Elderly residents’ family members arranging nearby assistance.")

add_heading_custom("4.3 User Personas (Summary)", 2)
add_bullet("Nimali (24, student): Sells used textbooks, borrows a kettle, finds a nearby math tutor.")
add_bullet("Kasun (33, resident): Needs a plumber urgently and wants neighborhood-verified options.")
add_bullet("Amaya (29, home baker): Lists homemade cakes and accepts local pickup orders.")
add_bullet(
    "Mr. Silva (68): Family posts a request for weekly grocery assistance from trusted neighbors."
)

add_heading_custom("5. Feasibility Study", 1)
add_heading_custom("5.1 Technical Feasibility", 2)
add_para(
    "Fully feasible within the assignment scope using Kotlin and Android Studio. Core "
    "features (authentication, listings, requests, categories, search/filter, ratings, and "
    "local database CRUD) can be implemented with standard Android components and a "
    "local/offline-capable database such as Room (SQLite). UI can be prototyped in Figma "
    "and built with XML or Jetpack Compose layouts. No dependency on unavailable proprietary "
    "APIs is required for MVP."
)

add_heading_custom("5.2 Operational Feasibility", 2)
add_para(
    "Users already understand posting, browsing, and messaging patterns from social apps. "
    "NeighbrLink simplifies this into clear modules: Community Help, Services, and "
    "Marketplace. Moderation can begin with basic report/block controls and admin flags for "
    "academic demonstration."
)

add_heading_custom("5.3 Economic Feasibility", 2)
add_para(
    "Development cost is low for a student MVP (no paid cloud requirement if local DB is used). "
    "Future monetization could include optional featured listings or small service boosting "
    "fees, but the academic version remains free and self-contained."
)

add_heading_custom("5.4 Schedule Feasibility", 2)
add_para(
    "Aligned with SDLC stages required by the module: proposal approval → requirements → "
    "Figma UI → Kotlin development with GitHub commits → unit/integration testing → final "
    "documentation. Scope is intentionally MVP-focused to remain deliverable."
)

add_heading_custom("6. Scope of the Application", 1)
add_heading_custom("6.1 In Scope (MVP)", 2)
add_bullet("User registration / login and profile management.")
add_bullet("Create, view, update, and delete community help posts (request/offer).")
add_bullet("Local marketplace listings for goods (sell / donate / borrow).")
add_bullet("Service provider profiles and service listings by category.")
add_bullet("Search and filter by category, keyword, and locality/area tag.")
add_bullet("Favorites / saved items.")
add_bullet("Basic rating and review for completed help or services.")
add_bullet("Simple inquiry/contact request with status tracking.")
add_bullet("Report listing / user for safety.")
add_bullet("Local database persistence for all core entities.")

add_heading_custom("6.2 Out of Scope (Future / Not in Assignment MVP)", 2)
add_bullet("Real payment gateway / in-app checkout.")
add_bullet("Live GPS tracking of service providers.")
add_bullet("Advanced AI matching or recommendation engines.")
add_bullet("iOS version.")
add_bullet(
    "Deployment to Play Store production release and long-term maintenance (excluded by assignment)."
)
add_bullet("Complex multi-admin enterprise moderation dashboards.")

add_heading_custom("7. Objectives", 1)
add_bullet("Provide a trusted hyperlocal digital space for community support and local exchange.")
add_bullet("Reduce dependence on unstructured social media groups for urgent neighborhood needs.")
add_bullet("Help local micro-service providers gain visible, category-based discovery.")
add_bullet("Enable safer buy/sell/borrow interactions within a defined community area.")
add_bullet("Demonstrate complete SDLC execution from planning to testing using Kotlin.")

add_heading_custom("8. Features and Functionalities", 1)
add_heading_custom("8.1 Core Modules", 2)
add_bullet("Auth & Profile: Sign up, login, edit profile, area/community selection, skills/tags.")
add_bullet(
    "Community Help Board: Post requests (blood, volunteers, lost & found, elder help) and "
    "offers; mark as open/closed."
)
add_bullet(
    "Local Marketplace: List goods with images, price/free/borrow mode, condition, and area."
)
add_bullet(
    "Services Directory: Create service ads (tuition, repairs, cleaning, caregiving, homemade "
    "food) with availability."
)
add_bullet(
    "Discovery: Home feed combining recent help posts, services, and marketplace items near "
    "selected area."
)
add_bullet("Trust Layer: Ratings, reviews, verification badge (manual/demo), report/block.")
add_bullet("My Activity: Manage my posts, saved items, and response status.")

add_heading_custom("8.2 Key Differentiators", 2)
add_bullet("Three-in-one hyperlocal model: Help + Services + Marketplace.")
add_bullet("Community-area first design (not generic nationwide classifieds).")
add_bullet("Borrow/donate modes in addition to selling.")
add_bullet("Trust and safety controls designed for neighborhood use.")

add_heading_custom("9. Defining Requirements", 1)
add_heading_custom("9.1 Functional Requirements", 2)
add_bullet("FR01: The system shall allow users to register and authenticate securely.")
add_bullet(
    "FR02: The system shall allow users to create and manage profiles including locality and "
    "contact preferences."
)
add_bullet(
    "FR03: The system shall allow users to create community help requests and offers with "
    "category and urgency level."
)
add_bullet(
    "FR04: The system shall allow users to publish marketplace listings (sell/donate/borrow) "
    "with details and images."
)
add_bullet("FR05: The system shall allow service providers to publish and manage service listings.")
add_bullet(
    "FR06: The system shall allow users to browse, search, and filter posts by category, "
    "keyword, and area."
)
add_bullet("FR07: The system shall allow users to save/favorite listings.")
add_bullet("FR08: The system shall allow users to submit ratings and reviews after interactions.")
add_bullet(
    "FR09: The system shall allow users to update listing status (available, reserved, "
    "completed, closed)."
)
add_bullet("FR10: The system shall persist all user and listing data in a local database.")
add_bullet("FR11: The system shall allow users to report inappropriate content.")
add_bullet(
    "FR12: The system shall provide a dashboard/home screen summarizing nearby activity."
)

add_heading_custom("9.2 Non-Functional Requirements", 2)
add_bullet(
    "NFR01 – Usability: New users should complete core tasks (post/browse) with minimal "
    "learning time; clear navigation."
)
add_bullet(
    "NFR02 – Performance: Common screens (feed, search results) should load smoothly on "
    "mid-range Android devices."
)
add_bullet(
    "NFR03 – Reliability: CRUD operations must consistently persist and retrieve data without loss."
)
add_bullet(
    "NFR04 – Security: Passwords stored using safe hashing; sensitive actions protected by "
    "login session."
)
add_bullet(
    "NFR05 – Maintainability: Modular Kotlin architecture (UI / ViewModel / Repository / DB layers)."
)
add_bullet(
    "NFR06 – Compatibility: Support modern Android API levels used in Android Studio target "
    "devices/emulators."
)
add_bullet(
    "NFR07 – Accessibility: Readable typography, adequate contrast, and touch-friendly controls."
)
add_bullet(
    "NFR08 – Scalability (design-level): Data model should allow future migration to cloud backend."
)

add_heading_custom("10. High-Level Design Direction (Preview)", 1)
add_heading_custom("10.1 Suggested Information Architecture", 2)
add_bullet("Splash / Onboarding")
add_bullet("Login & Register")
add_bullet("Home Feed (Help | Services | Marketplace tabs or segmented view)")
add_bullet("Create Post (type selector)")
add_bullet("Listing Details")
add_bullet("Provider / User Profile")
add_bullet("My Posts & Favorites")
add_bullet("Ratings & Reports")
add_bullet("Settings / Logout")

add_heading_custom("10.2 Design Notes for Figma Phase", 2)
add_bullet(
    "Apply 60-30-10 color rule (example direction: 60% soft neutral background, 30% deep "
    "teal/forest community brand, 10% warm amber CTA accent)."
)
add_bullet("Prioritize clarity over visual clutter; one primary action per screen.")
add_bullet(
    "Use consistent list patterns for help posts, services, and goods while keeping module "
    "identity clear."
)
add_bullet("Mobile-first wireframes → high-fidelity prototype before Kotlin UI implementation.")

add_heading_custom("10.3 Proposed Tech Stack", 2)
add_bullet("Language: Kotlin")
add_bullet("IDE: Android Studio")
add_bullet("UI: XML Layouts and/or Jetpack Compose")
add_bullet("Architecture: MVVM")
add_bullet("Database: Room (SQLite)")
add_bullet("Version Control: GitHub with regular commits from day one")
add_bullet("UI Prototype: Figma")

add_heading_custom("11. SDLC Plan (Assignment Alignment)", 1)
add_bullet("Planning & Requirement Analysis: Problem, audience, feasibility, scope (this proposal).")
add_bullet("Defining Requirements: Functional/non-functional requirements and feature set.")
add_bullet("Design: Complete Figma UI prototype with 60-30-10 color application.")
add_bullet(
    "Development: Kotlin implementation of layouts, interactivity, core features, and DB integration."
)
add_bullet("Testing: Unit testing, integration testing, and documented results.")
add_bullet(
    "Documentation: Final report with overview, features, challenges, screenshots, testing, "
    "and public GitHub link."
)
add_bullet("Excluded by module: Deployment and Maintenance stages.")

add_heading_custom("12. Risks & Mitigation", 1)
add_bullet("Scope creep → Strict MVP boundaries; freeze nice-to-have features until core CRUD flows work.")
add_bullet(
    "Trust/safety concerns in demo → Include report/block and status fields even if moderation "
    "is simplified."
)
add_bullet("UI complexity across 3 modules → Shared components and consistent navigation pattern.")
add_bullet(
    "Time pressure near testing → Build database and core screens early; test continuously with "
    "GitHub milestones."
)

add_heading_custom("13. Success Criteria", 1)
add_bullet("Panel approval of the idea under Topic 8.")
add_bullet("Complete Figma prototype covering all main user journeys.")
add_bullet(
    "Working Kotlin app with database-backed community help, services, and marketplace modules."
)
add_bullet("Evidence of regular GitHub commits.")
add_bullet("Documented unit and integration testing results.")

add_heading_custom("14. Conclusion", 1)
add_para(
    "NeighbrLink is a practical, innovative, and academically achievable response to the "
    "assigned topic. It solves a clear local problem—fragmented community help and low-trust "
    "neighborhood trading—by combining community services and a local marketplace in one Kotlin "
    "Android application. The idea has a visible market gap, a well-defined MVP scope, and strong "
    "alignment with the SE4041 marking criteria for ideation, design, development, database use, "
    "and documentation."
)
add_para(
    "Request to Panel: Approval is requested to proceed with Figma UI design and Kotlin "
    "development of NeighbrLink under Topic 8 – Community Services and Local Marketplace "
    "Mobile Application.",
    bold=True,
)

doc.add_paragraph()
add_heading_custom("Appendix A – Alternative Ideas Considered", 1)
add_para("The following alternatives were evaluated before selecting NeighbrLink:")
add_bullet(
    "CampusSwap Only: University-only buy/sell + tutoring. Strong for students, but narrower "
    "community impact."
)
add_bullet(
    "VolunteerBlood Connect: Community service focused on donations/volunteering only. Weaker "
    "marketplace coverage for Topic 8."
)
add_bullet(
    "ServiceBook LK: Local service booking only (plumber/tutor). Missing goods marketplace and "
    "mutual-aid depth."
)
add_para(
    "NeighbrLink was selected because it best balances innovation, market gap, topic coverage, "
    "and MVP feasibility."
)

add_heading_custom("Appendix B – Sample Categories", 1)
add_para(
    "Community Help: Blood Request, Volunteer Needed, Lost & Found, Elder Assistance, "
    "Education Help, Emergency Neighbor Support"
)
add_para(
    "Services: Home Repair, Tuition, Cleaning, Caregiving, Homemade Food, Beauty/Grooming at "
    "Home, Moving Help"
)
add_para(
    "Marketplace: Electronics, Books, Furniture, Fashion, Kitchen Items, Free/Donate, Borrow Tools"
)

out = r"D:\MAD\Assigment 01\NeighbrLink_Phase01_Proposal.docx"
doc.save(out)
print("Saved:", out)
